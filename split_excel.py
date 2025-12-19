#!/usr/bin/env python
"""
将包含多个工作表的Excel文件拆分成单独的Excel文件
"""

import os
import sys
from openpyxl import load_workbook, Workbook

# 添加项目根目录到Python路径
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

def split_excel_file():
    # 源文件路径
    source_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "testdata", "模拟数据.xlsx")
    
    # 检查源文件是否存在
    if not os.path.exists(source_file):
        print(f"❌ 源文件不存在：{source_file}")
        return
    
    # 加载源文件
    print(f"📂 加载源文件：{source_file}")
    wb = load_workbook(source_file)
    
    # 获取所有工作表名称
    sheet_names = wb.sheetnames
    print(f"📋 包含工作表：{sheet_names}")
    
    # 为每个工作表创建一个新的Excel文件
    for sheet_name in sheet_names:
        # 创建新工作簿
        new_wb = Workbook()
        
        # 获取新工作簿的活动工作表
        new_sheet = new_wb.active
        new_sheet.title = sheet_name
        
        # 获取源工作表
        source_sheet = wb[sheet_name]
        
        # 复制所有行
        for row in source_sheet.iter_rows(values_only=True):
            new_sheet.append(row)
        
        # 新文件路径
        new_file_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "testdata", f"{sheet_name}.xlsx")
        
        # 保存新文件
        new_wb.save(new_file_path)
        print(f"✅ 已生成：{new_file_path}")
    
    print("\n🎉 拆分完成！")

if __name__ == "__main__":
    split_excel_file()
