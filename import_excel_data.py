#!/usr/bin/env python
"""
将Excel文件中的社区工作者和干部数据导入到数据库
"""

import os
import sys
from datetime import datetime

# 添加项目根目录到Python路径
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

# 设置Django环境变量
os.environ.setdefault("DJANGO_SETTINGS_MODULE",
                      "community_management.settings")

import django

django.setup()

from openpyxl import load_workbook
from django.contrib.auth.models import User
from api.models import Community, SocialWorker, UserProfile, Cadre


def import_social_worker_data():
    """导入社区工作者数据"""
    print("开始导入社区工作者数据...")

    # Excel文件路径
    excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                              "testdata", "社区工作者.xlsx")

    # 加载Excel文件
    wb = load_workbook(excel_path)
    sheet = wb.active

    # 跳过表头，从第2行开始读取数据
    row_count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        # 解析行数据
        seq, name, gender, phone, birth_date, community_name, is_secretary, is_baned, remark = row

        # 处理社区
        community, _ = Community.objects.get_or_create(name=community_name)

        # 处理性别
        gender_map = {"男": 1, "女": 2}
        gender_value = gender_map.get(gender, 1)

        # 处理是否为书记
        is_secretary_value = True if is_secretary == "是" else False

        # 处理是否在职
        is_baned_value = True if is_baned == "在职" else False

        # 转换出生年月为date对象
        birth_date_obj = datetime.strptime(birth_date, "%Y-%m-%d").date()

        # 创建或更新社区工作者
        social_worker, created = SocialWorker.objects.update_or_create(
            name=name,
            defaults={
                "community": community,
                "gender": gender_value,
                "phone": phone,
                "birth_date": birth_date_obj,
                "is_secretary": is_secretary_value,
                "is_baned": is_baned_value,
                "remark": remark
            })

        row_count += 1
        action = "创建" if created else "更新"
        print(f"{action}社区工作者：{name} ({community_name})")

    print(f"\n社区工作者数据导入完成，共处理{row_count}条记录\n")
    return row_count


def import_cadre_data():
    """导入干部数据"""
    print("开始导入干部数据...")

    # Excel文件路径
    excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                              "testdata", "干部.xlsx")

    # 加载Excel文件
    wb = load_workbook(excel_path)
    sheet = wb.active

    # 跳过表头，从第2行开始读取数据
    row_count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        # 解析行数据
        seq, name, gender, phone, birth_date, position, is_baned, remark = row

        # 处理性别
        gender_map = {"男": 1, "女": 2}
        gender_value = gender_map.get(gender, 1)

        # 处理是否在职
        is_baned_value = True if is_baned == "在职" else False

        # 转换出生年月为date对象
        birth_date_obj = datetime.strptime(birth_date, "%Y-%m-%d").date()

        # 创建或更新干部
        cadre, created = Cadre.objects.update_or_create(name=name,
                                                        defaults={
                                                            "gender":
                                                            gender_value,
                                                            "phone": phone,
                                                            "birth_date":
                                                            birth_date_obj,
                                                            "position":
                                                            position,
                                                            "is_baned":
                                                            is_baned_value,
                                                            "remark": remark
                                                        })

        row_count += 1
        action = "创建" if created else "更新"
        print(f"{action}干部：{name} ({position})")

    print(f"\n干部数据导入完成，共处理{row_count}条记录\n")
    return row_count


def main():
    """主函数"""
    print("=============================================")
    print("开始导入Excel数据到数据库")
    print("=============================================")

    # 导入社区工作者数据
    social_worker_count = import_social_worker_data()

    # 导入干部数据
    cadre_count = import_cadre_data()

    print("=============================================")
    print("数据导入完成！")
    print(f"总处理记录数：")
    print(f"  - 社区工作者：{social_worker_count}条")
    print(f"  - 干部：{cadre_count}条")
    print(f"  - 总计：{social_worker_count + cadre_count}条")
    print("=============================================")


if __name__ == "__main__":
    main()
